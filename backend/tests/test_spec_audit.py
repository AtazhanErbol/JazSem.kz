import io
from decimal import Decimal

import pytest
from django.core.cache import cache
from django.db import connection
from django.test.utils import CaptureQueriesContext
from rest_framework.test import APIClient

from apps.ai.extraction import extract_pages
from apps.ai.models import AICourseDraft, AIJob, DocumentChunk, SourceDocument
from apps.ai.services import import_draft
from apps.assignments.models import Assignment
from apps.assignments.services import review, submit
from apps.courses.models import Course
from apps.courses.services import duplicate, new_version
from apps.enrollments.models import Enrollment
from apps.grading.services import grades
from apps.testing.models import Test as Quiz
from apps.testing.services import start


@pytest.mark.parametrize("role", ["admin", "teacher", "student"])
def test_all_read_endpoints_return_structured_responses(world, client_for, role):
    client = client_for(world[role])
    for endpoint in [
        "users",
        "disciplines",
        "groups",
        "courses",
        "assignments",
        "submissions",
        "tests",
        "attempts",
        "enrollments",
        "notifications",
        "audit",
        "content",
        "sources",
        "ai-jobs",
        "ai-usage",
    ]:
        result = client.get(f"/api/v1/{endpoint}/")
        assert result.status_code in [200, 403], (role, endpoint, result.data)
        assert "application/json" in result["Content-Type"]
        if result.status_code == 200:
            assert {"count", "next", "previous", "results"} <= result.data.keys()
    assert client.get("/api/v1/dashboard/").status_code == 200


def test_result_pagination_and_display_names(world, client_for):
    client = client_for(world["teacher"])
    data = client.get("/api/v1/enrollments/").data["results"][0]
    assert data["course_title"] == world["course"].title
    assert data["student_name"] == world["student"].email
    with CaptureQueriesContext(connection) as queries:
        client.get("/api/v1/enrollments/")
    assert len(queries) < 8
    for i in range(26):
        course = Course.objects.create(
            discipline=world["course"].discipline, teacher=world["teacher"], title=f"Page {i}"
        )
        Enrollment.objects.create(
            course=course,
            course_version=new_version(course, world["teacher"]),
            student=world["student"],
            assigned_by=world["teacher"],
        )
    first = client.get("/api/v1/enrollments/").data
    second = client.get("/api/v1/enrollments/?page=2").data
    assert first["count"] == 27 and len(first["results"]) == 25
    assert first["next"] and len(second["results"]) == 2


def test_assignment_deadline_ordering(world, client_for):
    from datetime import timedelta

    from django.utils import timezone

    world["assignment"].deadline = timezone.now() + timedelta(days=2)
    world["assignment"].save()
    earlier = Assignment.objects.create(
        topic=world["topic"],
        title="Earlier",
        instructions="x",
        status="PUBLISHED",
        deadline=timezone.now() + timedelta(days=1),
    )
    rows = (
        client_for(world["student"]).get("/api/v1/assignments/?ordering=deadline").data["results"]
    )
    assert rows[0]["id"] == str(earlier.pk)


def test_grading_query_count_does_not_grow_per_assignment(world):
    submission = submit(world["assignment"], world["student"], "answer", [])
    review(submission, world["teacher"], "grade", score=80)
    enrollment = Enrollment.objects.select_related("course_version").get(pk=world["enrollment"].pk)
    with CaptureQueriesContext(connection) as initial:
        assert grades(enrollment)["score"] == 40
    Assignment.objects.bulk_create(
        [Assignment(topic=world["topic"], title=f"Load {i}", instructions="x") for i in range(60)]
    )
    with CaptureQueriesContext(connection) as expanded:
        result = grades(enrollment)
    assert len(expanded) == len(initial) <= 6
    assert result["score"] == round(80 / 61 * 0.5, 2)


def test_attempt_filter_does_not_mix_tests_or_students(world, client_for):
    attempt = start(world["test"], world["student"])
    other_test = Quiz.objects.create(topic=world["topic"], title="Second quiz", status="PUBLISHED")
    response = client_for(world["student"]).get(f"/api/v1/attempts/?test={other_test.pk}")
    assert response.data["count"] == 0
    response = client_for(world["student"]).get(f"/api/v1/attempts/?test={world['test'].pk}")
    assert response.data["results"][0]["id"] == str(attempt.pk)
    assert client_for(world["outsider"]).get(f"/api/v1/attempts/{attempt.pk}/").status_code == 404


def test_submission_exposes_grading_scale_without_answer_keys(world, client_for):
    world["assignment"].max_score = 20
    world["assignment"].save()
    submission = submit(world["assignment"], world["student"], "answer", [])
    data = client_for(world["teacher"]).get(f"/api/v1/submissions/{submission.pk}/").data
    assert data["max_score"] == 20 and data["assignment_title"] == world["assignment"].title
    review(submission, world["teacher"], "grade", score=15)
    submission.refresh_from_db()
    assert submission.score == Decimal("75")


def test_ai_import_keeps_reviewed_title_description(world):
    source = SourceDocument.objects.create(
        course=world["course"],
        uploaded_by=world["teacher"],
        filename="x.txt",
        size=1,
        mime_type="text/plain",
    )
    chunk = DocumentChunk.objects.create(document=source, page_number=1, chunk_index=0, content="x")
    job = AIJob.objects.create(course=world["course"], user=world["teacher"])
    data = {
        "title": "Reviewed title",
        "description": "Reviewed description",
        "source_gaps": [],
        "weeks": [
            {
                "title": "Week",
                "topics": [
                    {
                        "title": "Topic",
                        "content": "x",
                        "source_chunks": [str(chunk.pk)],
                        "assignments": [],
                        "questions": [],
                    }
                ],
            }
        ],
    }
    draft = AICourseDraft.objects.create(job=job, data=data)
    version = import_draft(draft, world["teacher"])
    world["course"].refresh_from_db()
    assert world["course"].title == "Reviewed title"
    assert world["course"].description == "Reviewed description"
    assert version.status == "DRAFT"
    assert import_draft(draft, world["teacher"]).pk == version.pk


def test_blocked_user_loses_existing_session(world, client_for):
    cache.clear()
    client = APIClient()
    assert (
        client.post(
            "/api/v1/auth/login/",
            {"email": world["student"].email, "password": "Example-pass-583!"},
        ).status_code
        == 200
    )
    assert (
        client_for(world["teacher"])
        .post(f"/api/v1/users/{world['student'].pk}/deactivate/")
        .status_code
        == 204
    )
    assert client.get("/api/v1/courses/").status_code == 403


def test_private_source_and_chunk_unavailable_to_student(world, client_for):
    source = SourceDocument.objects.create(
        course=world["course"],
        uploaded_by=world["teacher"],
        filename="x.txt",
        size=1,
        mime_type="text/plain",
    )
    chunk = DocumentChunk.objects.create(
        document=source, page_number=1, chunk_index=0, content="private source"
    )
    student = client_for(world["student"])
    assert student.get("/api/v1/chunks/").data["count"] == 0
    assert student.get(f"/api/v1/chunks/{chunk.pk}/").status_code == 404
    assert student.get(f"/api/v1/sources/{source.pk}/download/").status_code == 404


def test_duplicate_does_not_mutate_original_or_copy_student_answers(world):
    target = duplicate(world["version"], world["teacher"])
    assert target.pk != world["version"].pk and target.status == "DRAFT"
    assert target.weeks.first().topics.first().tests.first().questions.count() == 1
    assert world["enrollment"].course_version_id == world["version"].pk


def test_docx_pptx_txt_extraction():
    from docx import Document
    from pptx import Presentation

    doc = Document()
    doc.add_paragraph("Kazakh: Қазақша. Русский текст.")
    buffer = io.BytesIO()
    doc.save(buffer)
    assert "Қазақша" in list(extract_pages(buffer.getvalue(), "lecture.docx"))[0][1]
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Lecture"
    buffer = io.BytesIO()
    deck.save(buffer)
    assert "Lecture" in list(extract_pages(buffer.getvalue(), "lecture.pptx"))[0][1]
    assert list(extract_pages("Қазақша".encode(), "lecture.txt")) == [(1, "Қазақша")]
