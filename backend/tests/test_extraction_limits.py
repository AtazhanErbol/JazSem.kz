import io
from unittest.mock import patch

import pytest
from pypdf import PdfWriter

from apps.ai.extraction import extract_pages


def pdf(pages):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=100, height=100)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


@pytest.mark.parametrize(
    "data", [b"%PDF-broken", pdf(0), pdf(301)], ids=["malformed", "empty", "too-many-pages"]
)
def test_invalid_empty_oversized_pdf_is_rejected_before_ocr(data):
    with patch("pytesseract.image_to_string") as ocr:
        with pytest.raises(Exception):
            list(extract_pages(data, "source.pdf"))
    ocr.assert_not_called()


def test_pdf_ocr_opens_document_once_and_preserves_languages(settings):
    import pypdfium2 as pdfium

    text = "Қазақша әғқңөұүһі. Русский текст. English text."
    with (
        patch("pypdfium2.PdfDocument", wraps=pdfium.PdfDocument) as opened,
        patch("pytesseract.image_to_string", return_value=text) as ocr,
    ):
        result = list(extract_pages(pdf(2), "source.pdf"))
    assert result == [(1, text), (2, text)]
    assert opened.call_count == 1 and ocr.call_count == 2
    assert ocr.call_args.kwargs["lang"] == settings.OCR_LANGUAGES


def test_docx_pptx_and_utf8_preserve_kazakh_text():
    from docx import Document
    from pptx import Presentation

    text = "ӘҒҚҢӨҰҮҺІ әғқңөұүһі Русский English"
    doc = Document()
    doc.add_paragraph(text)
    data = io.BytesIO()
    doc.save(data)
    assert list(extract_pages(data.getvalue(), "source.docx")) == [(1, text)]
    slides = Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(0, 0, 100, 100).text = text
    data = io.BytesIO()
    slides.save(data)
    assert list(extract_pages(data.getvalue(), "source.pptx")) == [(1, text)]
    assert list(extract_pages(text.encode(), "source.txt")) == [(1, text)]


def test_image_pixel_limit_precedes_ocr(monkeypatch):
    from PIL import Image

    monkeypatch.setattr("apps.ai.extraction.MAX_PIXELS", 10)
    with Image.new("RGB", (4, 4)) as picture:
        data = io.BytesIO()
        picture.save(data, format="PNG")
    with patch("pytesseract.image_to_string") as ocr, pytest.raises(ValueError):
        list(extract_pages(data.getvalue(), "source.png"))
    ocr.assert_not_called()
